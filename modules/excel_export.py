from openpyxl import Workbook
import os

def exportar_a_excel(datos, perfil, plan_armado, plan_corte, piezas_usadas_inventario, output_folder, grafico_corte=None, grafico_armado=None):
    """
    Exporta los resultados a un archivo Excel con múltiples hojas.
    """
    # Sanitizar el nombre del perfil para usarlo en el nombre del archivo
    nombre_perfil_sanitizado = "".join(c for c in perfil if c.isalnum() or c in " _-")
    filename = f"Reporte_Corte_{nombre_perfil_sanitizado}.xlsx"
    filepath = os.path.join(output_folder, filename)

    # Crear libro de Excel
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Resumen"

    # === Hoja 1: Resumen ===
    ws1.append(["Perfil", perfil])
    ws1.append(["Total de barras usadas", len(datos['barras'])])
    ws1.append(["Piezas usadas del inventario", len(piezas_usadas_inventario)])
    ws1.append(["Penalización por partes chicas", datos['penalizacion_partes_chicas']])
    ws1.append([""])

    # === Hoja 2: Plan de Corte ===
    ws2 = wb.create_sheet(title="Plan de Corte")
    ws2.append(["ID Barra", "Partes", "Desperdicio (mm)"])
    # === Agregar barras de inventario ===
    for i, longitud in enumerate(piezas_usadas_inventario):
        id_barra = f"BAR-{i + 1:03d}"
        ws2.append([id_barra, f"INV-{i+1:03d}({longitud})", 0])
    # === Agregar barras nuevas ===
    for barra in datos['barras']:
        partes_str = [f"{p[0]}({p[1]})" for p in barra['partes']]  # Convertir tuplas a strings
        ws2.append([barra['id'], ", ".join(partes_str), barra['resto']])

    # === Hoja 3: Piezas del Inventario Usadas ===
    ws3 = wb.create_sheet(title="Inventario Usado")
    ws3.append(["Pieza", "Longitud (mm)"])
    for p in piezas_usadas_inventario:
        ws3.append(["Inventario", p])

    # === Hoja 4: Plan de Armado ===
    ws4 = wb.create_sheet(title="Plan de Armado")
    ws4.append(["Pieza", "Partes"])
    for pieza_num in sorted(plan_armado.keys()):
        partes_pieza = plan_armado[pieza_num]
        if len(partes_pieza) == 1:
            nombre, longitud = partes_pieza[0]
            ws4.append([f"Pieza-{pieza_num}", f"{nombre}({longitud})"])
        else:
            partes_str = [f"{nombre}({longitud})" for nombre, longitud in partes_pieza]
            ws4.append([f"Pieza-{pieza_num}", " + ".join(partes_str)])

    # === Hoja 5: Resultado Tabla ===
    ws5 = wb.create_sheet(title="Resultado Tabla")
    ws5.append(["Pieza", "Partes"])
    for pieza_num in sorted(plan_armado.keys()):
        etiquetas = [f"{nombre}({longitud})" for nombre, longitud in plan_armado[pieza_num]]
        ws5.append([f"Pieza-{pieza_num}", "\t".join(etiquetas)])

    # === No insertar gráficos ===

    # Guardar archivo
    wb.save(filepath)
    print(f"\n✅ Reporte exportado a: {filepath}")

    return filename